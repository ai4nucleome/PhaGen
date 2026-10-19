#!/usr/bin/env python3


from __future__ import annotations

import argparse
import openpyxl
import numpy as np
import pandas as pd
from Bio import SeqIO
from tqdm import tqdm
from pathlib import Path
from t7_gene_coords import T7_GENE_COORDS, T7_GENE_DICT


def load_sd01(path, gene_coord, gene_dict):
    
    wb = openpyxl.load_workbook(path)
    records = []

    for gene in wb.sheetnames:
        if gene in ["19.3"]:
            continue

        coords = gene_coord.get(gene_dict.get(gene, gene), None)
        if coords is None:
            print(f"Missing coordination: {gene}")
            continue

        ws = wb[gene]
        rows = list(ws.iter_rows(values_only=True))

        assert (
            rows[0][17] == "Average NMI Rep A" and rows[0][20] == "Average NMI Rep B"
        ), gene

        for row in rows[1:]:
            if row[0] is None or row[1] is None:
                continue

            mut_str = str(row[0])
            if (
                "," not in mut_str
                or int(row[1]) != int(mut_str.split(",")[0][1:-1])
                or (int(row[1]) < coords[0] or int(row[1]) > coords[1])
            ):
                continue

            t7_pos = int(row[1])
            wt_nt = mut_str.split(",")[0][0]
            mut_nt = mut_str.split(",")[0][-1]

            try:
                nmi_a = float(row[17]) if row[17] is not None else None
                nmi_b = float(row[20]) if row[20] is not None else None
            except (TypeError, ValueError) as e:
                raise ValueError(
                    f"NMI 列无法解析 (gene={gene!r}, repA={row[17]!r}, repB={row[20]!r})"
                ) from e

            nmis = [x for x in [nmi_a, nmi_b] if x is not None]
            avg_nmi = float(np.mean(nmis)) if nmis else None

            aa_part = mut_str.split(",")[1] if "," in mut_str else ""
            if "*" in aa_part:
                mut_type = "stop"
            elif len(aa_part) >= 2 and aa_part[0] == aa_part[-1]:
                mut_type = "synonymous"
            else:
                mut_type = "nonsynonymous"

            records.append(
                {
                    "gene": gene,
                    "mutation": mut_str,
                    "wt_nt": wt_nt,
                    "nt_position": t7_pos,
                    "mut_nt": mut_nt,
                    "mut_type": mut_type,
                    "nmi_repA": nmi_a,
                    "nmi_repB": nmi_b,
                    "avg_nmi": avg_nmi,
                }
            )

    wb.close()
    return pd.DataFrame(records)


def _build_mutant_rows(
    gene_sheet: str, gene_df: pd.DataFrame, genome: str, gene_coord
) -> list[dict]:
    wt_cds = genome[gene_coord[0] - 1 : gene_coord[1]]
    out: list[dict] = []

    for _, row in gene_df.iterrows():
        nmi = row["avg_nmi"]
        if nmi is None or (isinstance(nmi, float) and pd.isna(nmi)):
            continue

        mut_str = str(row["mutation"])
        abs_pos = row["nt_position"]
        ref_nt = row["wt_nt"]
        alt_nt = row["mut_nt"]
        assert genome[abs_pos - 1].upper() == ref_nt, "WT is wrong"

        rel_pos = abs_pos - gene_coord[0]
        assert 0 <= rel_pos < len(wt_cds), "Wrong"

        mutant_cds = wt_cds[:rel_pos] + alt_nt + wt_cds[rel_pos + 1 :]
        out.append(
            {
                "gene": gene_sheet,
                "seq": mutant_cds,
                "nmi": float(nmi),
                "mutation": mut_str,
                "mut_type": row["mut_type"],
            }
        )
    return out


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Prepare T7 mutant CDS TSVs from sd01.xlsx"
    )
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument(
        "--data_dir",
        type=Path,
        default=default_root / "data",
        help="Directory containing robins/sd01.xlsx and T7/V01146.1.fasta",
    )
    parser.add_argument(
        "--out_dir",
        type=Path,
        default=None,
        help="Output directory (default: <data_dir>/processed)",
    )
    parser.add_argument(
        "--force", action="store_true", help="Replace existing processed TSV files."
    )
    args = parser.parse_args()

    data_dir = args.data_dir.resolve()
    out_dir = (args.out_dir or (data_dir / "processed")).resolve()
    existing = sorted(out_dir.glob("gene_*.tsv")) + (
        [out_dir / "all_genes.tsv"] if (out_dir / "all_genes.tsv").exists() else []
    )
    if existing and not args.force:
        raise FileExistsError(
            f"Output already exists: {existing[0]}; use --force to replace it."
        )
    out_dir.mkdir(parents=True, exist_ok=True)

    sd01_path = data_dir / "robins" / "sd01.xlsx"
    fasta_path = data_dir / "T7" / "V01146.1.fasta"
    if not fasta_path.is_file():
        raise FileNotFoundError(f"FASTA not found: {fasta_path}")

    record = SeqIO.read(fasta_path, "fasta")
    genome = str(record.seq).upper()
    if len(genome) < 39000:
        raise RuntimeError(
            f"Genome sequence unexpectedly short ({len(genome)} nt): {fasta_path}"
        )

    df = load_sd01(str(sd01_path), T7_GENE_COORDS, T7_GENE_DICT)

    all_parts: list[pd.DataFrame] = []
    per_gene_counts: list[tuple[str, int]] = []

    for gene in tqdm(df["gene"].unique()):
        gdf = df[df["gene"] == gene]
        gene_coord = T7_GENE_COORDS.get(T7_GENE_DICT.get(gene, gene), None)
        if gene_coord is None:
            print(f"Warning: no coords for gene {gene}, skipping")
            continue

        rows = _build_mutant_rows(gene, gdf, genome, gene_coord)
        gout = (
            pd.DataFrame(rows)
            if rows
            else pd.DataFrame(columns=["gene", "seq", "nmi", "mutation", "mut_type"])
        )
        gout.to_csv(out_dir / f"gene_{gene}.tsv", sep="\t", index=False)
        per_gene_counts.append((gene, len(gout)))
        all_parts.append(gout)

    non_empty = [p for p in all_parts if not p.empty]
    all_df = (
        pd.concat(non_empty, ignore_index=True)
        if non_empty
        else pd.DataFrame(columns=["gene", "seq", "nmi", "mutation", "mut_type"])
    )
    all_df.to_csv(out_dir / "all_genes.tsv", sep="\t", index=False)

    total = int(all_df.shape[0])
    print(f"Wrote {len(per_gene_counts)} gene TSVs + all_genes.tsv under {out_dir}")
    print(f"sd01: {sd01_path}")
    print(f"Total rows: {total}")
    print("Rows per gene:")
    for name, n in sorted(per_gene_counts, key=lambda x: (-x[1], x[0])):
        print(f"  {name}: {n}")


if __name__ == "__main__":
    main()
